"""Презентация Maqsat для финала: py docs/build_deck.py → docs/Maqsat_VentureHack2026_Final.pptx
Стиль как у команды на отборе: синий + чёрный, Arial Black, фон-тоннель; мотив — кольца мишени (логотип Maqsat).
"""
from pathlib import Path
import math
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = Path(__file__).parent
IMG = HERE / 'img'
OUT = HERE / 'Maqsat_VentureHack2026_Final.pptx'

BLUE = RGBColor(0x2F, 0x3F, 0xBF)
BLUE_DARK = RGBColor(0x1A, 0x24, 0x7A)
BLUE_MID = RGBColor(0x6E, 0x7B, 0xE0)
BLUE_LIGHT = RGBColor(0xC3, 0xC9, 0xF4)
BLUE_PALE = RGBColor(0xEC, 0xEE, 0xFC)
GOLD = RGBColor(0xF2, 0xA9, 0x00)
GOLD_PALE = RGBColor(0xFF, 0xF1, 0xD6)
INK = RGBColor(0x11, 0x14, 0x24)
GRAY = RGBColor(0x4A, 0x4F, 0x62)
MUTED = RGBColor(0x7A, 0x80, 0x94)
LINE = RGBColor(0xDA, 0xDE, 0xEC)
BG = RGBColor(0xF6, 0xF7, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x0C, 0x7A, 0x5C)
GREEN_PALE = RGBColor(0xDC, 0xF2, 0xEA)
ORANGE = RGBColor(0x9A, 0x52, 0x00)
ORANGE_PALE = RGBColor(0xFD, 0xEC, 0xD6)
VIOLET = RGBColor(0x5B, 0x48, 0xC9)
VIOLET_PALE = RGBColor(0xEB, 0xE7, 0xFB)
RED = RGBColor(0xB3, 0x31, 0x2F)
RED_PALE = RGBColor(0xFB, 0xE3, 0xE2)

FONT = 'Arial Black'
SCALE = 0.88          # Arial Black широкий — обычный текст чуть мельче
FOOTER = 'VentureHack 2026  ·  Финал  ·  EduTech'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5


def make_tunnel(path, w=2000, h=1125):
    img = Image.new('RGB', (w, h), (246, 247, 251))
    d = ImageDraw.Draw(img)
    for y in range(h):
        c = int(246 + 8 * y / h)
        d.line([(0, y), (w, y)], fill=(c, c + 1, 255))
    vx, vy = int(w * 0.66), int(h * 0.47)
    radii = [55 * 1.32 ** k for k in range(15)]
    for i, r in enumerate(radii):
        if i % 2 == 0 and i + 1 < len(radii):
            R = radii[i + 1]
            d.ellipse([vx - R, vy - R * 0.72, vx + R, vy + R * 0.72], fill=(240, 241, 252))
            d.ellipse([vx - r, vy - r * 0.72, vx + r, vy + r * 0.72], fill=(247, 248, 254))
    for r in radii:
        d.ellipse([vx - r, vy - r * 0.72, vx + r, vy + r * 0.72], outline=(214, 219, 240), width=3)
    for a in range(0, 360, 30):
        R = radii[-1]
        d.line([(vx, vy), (vx + R * math.cos(math.radians(a)), vy + R * 0.72 * math.sin(math.radians(a)))], fill=(226, 230, 246), width=2)
    img.save(path)


TUNNEL = IMG / '_tunnel.png'
make_tunnel(TUNNEL)


def bg(s, color):
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(s, x, y, w, h, runs, size=16, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if spacing:
            para.space_after = Pt(spacing)
        for t, o in (p if isinstance(p, list) else [(p, {})]):
            r = para.add_run()
            r.text = t
            sz = o.get('size', size)
            r.font.name = FONT
            r.font.size = Pt(sz if sz >= 22 else round(sz * SCALE))
            r.font.bold = False
            r.font.italic = o.get('italic', italic)
            r.font.color.rgb = o.get('color', color)
    return tb


def shape(s, kind, x, y, w, h, fill=None, line=None, lw=2, dash=False):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = 4
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    return sh


def circle(s, cx, cy, d, fill=None, line=None, lw=2):
    return shape(s, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill, line, lw)


def label(sh, t, size=14, color=WHITE, align=PP_ALIGN.CENTER):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(t.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.name = FONT
        r.font.size = Pt(size if size >= 22 else round(size * SCALE))
        r.font.color.rgb = color


def line(s, x1, y1, x2, y2, color=BLUE_LIGHT, width=1.5, arrow=False, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = 4
    if arrow:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))


def target(s, cx, cy, d, color=BLUE, core=GOLD):
    """Мотив Maqsat — мишень."""
    circle(s, cx, cy, d, None, color, max(2, d * 3))
    circle(s, cx, cy, d * 0.58, None, color, max(2, d * 3))
    circle(s, cx, cy, d * 0.2, core)


def brand(s, right=True):
    x = SW - 1.95 if right else 0.6
    target(s, x + 0.21, 0.63, 0.42)
    text(s, x + 0.52, 0.44, 1.5, 0.4, 'Maqsat', size=16, color=BLUE, anchor=MSO_ANCHOR.MIDDLE)


def footer(s):
    text(s, SW - 5.1, SH - 0.55, 4.5, 0.3, FOOTER, size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def heading(s, big, sub, x=0.7, y=0.95, w=11.9):
    text(s, x, y, w, 0.85, big, size=44, color=BLUE)
    text(s, x, y + 0.82, w, 0.9, sub, size=22, color=INK)


def button(s, x, y, w, t, h=0.5, fill=BLUE, color=WHITE):
    b = shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, fill)
    label(b, t, size=14, color=color)


def picture(s, path, x, y, w, crop_bottom=0.0):
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    if crop_bottom:
        h = pic.height
        pic.crop_bottom = crop_bottom
        pic.height = int(h * (1 - crop_bottom))
    pic.line.color.rgb = LINE
    pic.line.width = Pt(1)


def dots(s, pts, color=BLUE):
    for x, y, d in pts:
        circle(s, x, y, d, color)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def bullets(tb, indent=0.22):
    for p in tb.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', str(Inches(indent)))
        pPr.set('indent', str(-Inches(indent)))
        pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '•'}))


# ===================================================================== 1. Титул
s = prs.slides.add_slide(BLANK)
s.shapes.add_picture(str(TUNNEL), 0, 0, width=prs.slide_width, height=prs.slide_height)
dots(s, [(1.5, 0.35, 0.2), (0.55, 1.05, 0.2), (3.5, 1.0, 0.2), (4.3, 0.5, 0.18)])
brand(s)
# «мишень» с тремя кольцами: возможность → план → результат
cx, cy = 2.95, 4.1
for d, fill, col, t in [(4.6, BLUE_PALE, BLUE_LIGHT, None), (3.3, BLUE_LIGHT, BLUE_MID, None), (2.0, BLUE, BLUE, None)]:
    circle(s, cx, cy, d, fill, col, 3)
core = circle(s, cx, cy, 0.9, GOLD)
label(core, 'ЦЕЛЬ', size=12, color=INK)
text(s, cx - 2.1, cy - 2.2, 4.2, 0.3, 'ВОЗМОЖНОСТЬ', size=12, color=BLUE, align=PP_ALIGN.CENTER)
text(s, cx - 1.5, cy - 1.55, 3.0, 0.3, 'ПЛАН', size=12, color=BLUE_DARK, align=PP_ALIGN.CENTER)
text(s, cx - 1.0, cy - 0.85, 2.0, 0.3, 'ШАГИ', size=11, color=WHITE, align=PP_ALIGN.CENTER)
text(s, 6.45, 1.95, 6.5, 1.2, 'MAQSAT', size=66, color=BLUE)
text(s, 6.5, 3.05, 6.6, 0.6, 'ОТ ВОЗМОЖНОСТИ К РЕЗУЛЬТАТУ', size=24, color=INK)
text(s, 6.5, 3.85, 6.4, 0.7, 'Подбор образовательных возможностей и план назад от дедлайна', size=17, color=BLUE)
text(s, 6.5, 4.65, 6.1, 1.0, 'Олимпиады, стипендии, обменные и летние программы, бесплатные IT-школы — кто подходит, чего не хватает и что делать на этой неделе. Мақсат (каз.) — цель.', size=15, color=GRAY)
button(s, 6.5, 5.85, 3.3, 'Финал · EduTech')
text(s, SW - 5.1, SH - 0.55, 4.5, 0.3, 'VentureHack 2026', size=12, color=MUTED, align=PP_ALIGN.RIGHT)
notes(s, 'Maqsat — по-казахски «цель». Мы помогаем пройти путь от «я где-то слышал о программе» до «заявка отправлена вовремя».')

# ===================================================================== 2. Проблема
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
brand(s)
heading(s, 'ПРОБЛЕМА', 'ВОЗМОЖНОСТИ ЕСТЬ — ДОЙТИ ДО НИХ СЛОЖНО')
cards = [('Узнал после дедлайна', 'Информация разбросана по десяткам сайтов и каналов. О конкурсе узнают, когда регистрация уже закрыта.', RED),
         ('Не понял, подходит ли', '«Sophomore or junior (or equivalent)», «16–17 лет на 1 сентября года поступления», «IELTS 6.0, Writing 6.0» — условия написаны не для школьника.', ORANGE),
         ('Начал слишком поздно', 'Сертификат IELTS, рекомендации и эссе требуют недель. За 2–3 недели до дедлайна часть шагов уже не успеть.', VIOLET)]
for i, (h_, b, c) in enumerate(cards):
    x = 0.7 + i * 4.1
    shape(s, MSO_SHAPE.RECTANGLE, x, 3.0, 3.8, 3.35, BG, LINE, 1)
    num = circle(s, x + 0.72, 3.0, 0.8, c, WHITE, 3)
    label(num, str(i + 1), size=20)
    text(s, x + 0.3, 3.65, 3.2, 0.8, h_, size=19, color=c)
    text(s, x + 0.3, 4.45, 3.25, 1.8, b, size=14, color=GRAY)
text(s, 0.7, 6.6, 11.9, 0.4, 'Все три ситуации — из текста финального кейса. Ниже — что мы увидели сами, когда собирали данные.', size=12, color=MUTED)
footer(s)
notes(s, 'Кейс финала начинается ровно с этих трёх ситуаций. Наш продукт отвечает на каждую: подбор, проверка условий, план.')

# ===================================================================== 3. Исследование
s = prs.slides.add_slide(BLANK)
bg(s, BG)
brand(s)
heading(s, 'ИССЛЕДОВАНИЕ', '«НЕ ЗНАЛ ВОВРЕМЯ» — БАРЬЕР №1')
stats = [('2', 'из 13 — даты нового\nцикла опубликованы', GREEN), ('8', 'из 13 — только оценка\nпо прошлому году', ORANGE),
         ('2', 'из 13 — даты нет\nвообще', RED), ('1', 'из 13 — набор\nидёт постоянно', BLUE)]
for i, (n, l, c) in enumerate(stats):
    x = 1.35 + i * 2.95
    ring = circle(s, x + 0.25, 3.55, 2.0, WHITE, c, 5)
    label(ring, n, size=44, color=c)
    text(s, x - 0.9, 4.75, 2.3, 0.8, l, size=13, color=GRAY, align=PP_ALIGN.CENTER)
text(s, 0.7, 5.5, 11.9, 0.4, 'Мы сами собрали 13 возможностей для школьников Казахстана по официальным сайтам и открытым источникам 26.09.2026.', size=13, color=INK)
survey = [('59%', 'главный барьер —\n«не знал вовремя»'), ('5%', 'узнают на сайтах\nорганизаторов'),
          ('41%', 'не всегда понимают,\nподходят ли'), ('73%', 'воспользовались бы\n(да + возможно)')]
shape(s, MSO_SHAPE.RECTANGLE, 0.7, 5.95, 11.9, 1.15, WHITE, BLUE, 1.5)
text(s, 0.95, 6.1, 2.1, 0.85, [[('Опрос учеников', {'color': BLUE})], [('n = 22, 26.09.2026', {'color': MUTED, 'size': 11})]], size=14)
for i, (n, l) in enumerate(survey):
    x = 3.05 + i * 2.38
    text(s, x, 6.0, 2.25, 0.45, n, size=24, color=BLUE)
    text(s, x, 6.5, 2.25, 0.55, l, size=11, color=GRAY)
notes(s, 'Верхние цифры — наш собственный сбор данных (экран «Откуда данные»). Нижние — опрос учеников: 24 ответа, 2 исключены (бессмысленные), выборка смещена к 12 классу и студентам — так и говорите.')

# ===================================================================== 3б. Интервью
# Реальные результаты юзабилити-теста (секунды до первой цели в плане). Пусто — строка на слайде не выводится.
USABILITY = []
USABILITY_GRADES = '9–11 класс'
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
brand(s)
heading(s, 'ИНТЕРВЬЮ', '5 ЧЕЛОВЕК — И ЧТО МЫ ИЗМЕНИЛИ ЗА НОЧЬ')
quotes = [('«Иногда ученик узнаёт о программе за несколько дней до дедлайна и не успевает получить рекомендацию»', 'школьный профориентатор'),
          ('«Рекомендации и документы требуют участия школы — их нельзя подготовить в последний день»', 'заместитель директора'),
          ('«Не доверяю формулировке «полная стипендия» без ссылки на условия»', 'родитель')]
for i, (q, who) in enumerate(quotes):
    y = 2.55 + i * 1.3
    shape(s, MSO_SHAPE.RECTANGLE, 0.7, y, 6.3, 1.15, BG)
    text(s, 0.95, y + 0.12, 5.9, 0.7, q, size=13, color=INK, italic=True)
    text(s, 0.95, y + 0.82, 5.9, 0.3, '— ' + who, size=11, color=BLUE)
changes = [('🏫 через школу', 'шаги с рекомендациями и документами отмечены в плане'),
           ('Заявка отправлена', 'закрывает подготовку, этап отбора остаётся'),
           ('Свои задачи', 'в плане и календаре, даже без дедлайна'),
           ('Стоимость', 'в каждой карточке; «бесплатно» — только если подтверждено')]
text(s, 7.4, 2.45, 5.2, 0.4, 'ДОБАВИЛИ ПО ПРОСЬБАМ', size=12, color=MUTED)
for i, (h_, b) in enumerate(changes):
    y = 2.85 + i * 0.9
    target(s, 7.55, y + 0.2, 0.32)
    text(s, 7.9, y, 4.7, 0.4, h_, size=16, color=BLUE)
    text(s, 7.9, y + 0.38, 4.7, 0.5, b, size=12, color=GRAY)
if USABILITY:
    med = sorted(USABILITY)[len(USABILITY) // 2] if len(USABILITY) % 2 else sum(sorted(USABILITY)[len(USABILITY) // 2 - 1:len(USABILITY) // 2 + 1]) // 2
    mmss = lambda v: f'{v // 60}:{v % 60:02d}'
    text(s, 0.7, 6.65, 11.9, 0.4, f'Юзабилити-тест: {len(USABILITY)} {"ученика" if 2 <= len(USABILITY) <= 4 else "учеников"} ({USABILITY_GRADES}) сами нашли возможность и добавили её в план — медиана {mmss(med)} (от {mmss(min(USABILITY))} до {mmss(max(USABILITY))}).', size=12, color=BLUE)
else:
    text(s, 0.7, 6.65, 11.9, 0.4, 'Все пятеро доверяют только данным со ссылкой на источник и датой проверки — это уже было в продукте. «Шансы поступления» и шаблонные эссе нам прямо не советовали делать.', size=12, color=MUTED)
notes(s, 'Интервью: школьный профориентатор, завуч, учитель английского, ученик, подававший на FLEX, родитель. Полные ответы — docs/research-plan.md. Подчеркните: изменения сделаны в ту же ночь.')

# ===================================================================== 4. Решение
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
brand(s)
heading(s, 'РЕШЕНИЕ', 'ПРОФИЛЬ → ПОДБОР → ПЛАН → КАЛЕНДАРЬ')
steps = [('Профиль за минуту', 'Возраст, класс, английский или IELTS, средний балл, город, интересы. Без регистрации.'),
         ('Подбор с проверкой', 'Для каждой возможности: подходишь, почти (и чего не хватает) или станет доступно позже.'),
         ('План назад от дедлайна', 'Когда записаться на IELTS, просить рекомендации, писать эссе. И на годы вперёд: «доступно в 2028» → цель с планом с сегодняшнего дня.'),
         ('Календарь и прогресс', 'Все шаги — в календарь телефона с напоминаниями. «Сделай сегодня» и прогресс по целям.')]
for i, (h_, b) in enumerate(steps):
    x = 0.9 + i * 3.05
    if i < 3:
        line(s, x + 1.0, 3.0, x + 3.0, 3.0, BLUE_LIGHT, 2, dash=True)
    num = circle(s, x + 0.45, 3.0, 0.95, BLUE if i < 3 else GOLD, WHITE, 3)
    label(num, str(i + 1), size=22, color=WHITE if i < 3 else INK)
    text(s, x, 3.75, 2.8, 0.8, h_, size=19, color=BLUE)
    text(s, x, 4.55, 2.75, 1.9, b, size=14, color=GRAY)
button(s, 3.9, 6.35, 8.4, 'На любом телефоне и без интернета · без регистрации · русский и қазақша')
footer(s)

# ===================================================================== 5. Подбор
s = prs.slides.add_slide(BLANK)
bg(s, BG)
brand(s)
heading(s, 'ПОДБОР', 'НЕ СПИСОК, А ПРОВЕРКА УСЛОВИЙ')
rows = [('✓ ПОДХОДИШЬ', 'UWC для 10-классника', 'Возраст считается на 1 сентября года поступления: будет 17 — подходит.', GREEN, GREEN_PALE),
        ('△ ПОЧТИ', 'Yale Young Global Scholars', 'Английский B1, ориентир — B2. План сам добавит шаги по языку.', ORANGE, ORANGE_PALE),
        ('⏳ ПОЗЖЕ', 'Назарбаев Университет', 'Сейчас 10 класс → «доступно в 2028». Можно сразу сделать это целью: план начинается сегодня.', VIOLET, VIOLET_PALE),
        ('✗ НЕ ПОДХОДИТ', 'UWC без гражданства РК', 'Гражданство или ВНЖ — жёсткое условие; такие программы скрыты по умолчанию.', RED, RED_PALE)]
for i, (st, who, why, c, pale) in enumerate(rows):
    y = 2.6 + i * 0.95
    shape(s, MSO_SHAPE.RECTANGLE, 0.7, y, 11.9, 0.82, WHITE, LINE, 1)
    b = shape(s, MSO_SHAPE.RECTANGLE, 0.7, y, 2.6, 0.82, pale)
    label(b, st, size=15, color=c)
    text(s, 3.55, y + 0.1, 3.4, 0.62, who, size=15, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 7.0, y + 0.08, 5.45, 0.66, why, size=13, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.7, 6.55, 11.9, 0.4, 'Примеры — реальный результат приложения для демо-профиля: 10 класс, Алматы, английский B1.', size=12, color=MUTED)
notes(s, 'Ключевое отличие от каналов с объявлениями: мы не просто показываем возможность, а говорим, подходит ли она именно тебе и чего не хватает.')

# ===================================================================== 6. Демо: подбор
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
circle(s, 4.6, 4.5, 8.6, BLUE_PALE)
brand(s)
heading(s, 'ДЕМО', 'ВОЗМОЖНОСТИ ДЛЯ ТЕБЯ')
picture(s, IMG / 'matches.png', 0.7, 2.55, 7.3, crop_bottom=0.1)
pts = [('Статус и пробелы', 'подходишь / почти / позже — с объяснением по каждому условию'),
       ('Дедлайн и отсчёт', '«через 12 дней»; красным — если меньше 3 недель'),
       ('Честная дата', '«ориентировочно» и на чём основана оценка'),
       ('Источник', 'ссылка на официальный сайт и дата проверки')]
for i, (h_, b) in enumerate(pts):
    y = 2.5 + i * 1.0
    target(s, 8.7, y + 0.2, 0.34)
    text(s, 9.05, y, 3.9, 0.4, h_, size=18, color=BLUE)
    text(s, 9.05, y + 0.4, 3.8, 0.55, b, size=13, color=GRAY)
button(s, 9.05, 6.55, 3.4, 'Живое демо: #demo', h=0.45)
notes(s, 'Живое демо: главная → «Посмотреть на примере». Показать карточку ML School (дедлайн 8 октября), Yale (почти: английский), НУ (позже, 2028).')

# ===================================================================== 7. Демо: план
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
circle(s, 9.3, 4.5, 8.6, BLUE_PALE)
brand(s, right=False)
heading(s, 'ПЛАН', 'НАЗАД ОТ ДЕДЛАЙНА', w=5.0)
picture(s, IMG / 'plan.png', 5.3, 1.35, 7.35, crop_bottom=0.1)
pts = [('Сделай сегодня', 'один следующий шаг, а не список на полгода'),
       ('Конфликт дедлайнов', 'UWC 5 января и Yale 6 января — видно заранее'),
       ('Сжатый график', 'опоздал с началом — шаги сжимаются, есть более поздний раунд'),
       ('Календарь и печать', '.ics с напоминаниями; версия для родителей')]
for i, (h_, b) in enumerate(pts):
    y = 2.75 + i * 1.02
    target(s, 0.87, y + 0.2, 0.34)
    text(s, 1.25, y, 3.9, 0.4, h_, size=17, color=BLUE)
    text(s, 1.25, y + 0.38, 3.9, 0.55, b, size=12, color=GRAY)
notes(s, 'Живое демо: добавить 2–3 возможности → «Мой план» → показать предупреждение о двух дедлайнах → нажать «Добавить в календарь». Запасной путь: #demo-plan.')

# ===================================================================== 8. Честные данные
s = prs.slides.add_slide(BLANK)
bg(s, BG)
brand(s)
heading(s, 'ДАННЫЕ', 'МЫ НЕ ПРИДУМЫВАЕМ ДАТЫ')
levels = [('без пометки', 'Сверено с официальным сайтом организатора 26.09.2026', GREEN, GREEN_PALE),
          ('УТОЧНИТЕ', 'По открытым источникам или СМИ — проверьте перед подачей', ORANGE, ORANGE_PALE),
          ('НАША ОЦЕНКА', 'Организатор не называет значение (например, уровень английского)', VIOLET, VIOLET_PALE),
          ('ОРИЕНТИРОВОЧНО', 'Срок нового цикла не объявлен — дата по прошлому циклу и её основание', RGBColor(0xB8, 0x77, 0x00), GOLD_PALE)]
for i, (tag, desc, c, pale) in enumerate(levels):
    y = 2.6 + i * 0.9
    b = shape(s, MSO_SHAPE.RECTANGLE, 0.7, y, 3.1, 0.72, pale, c, 1.5)
    label(b, tag, size=14, color=c)
    text(s, 4.05, y + 0.1, 8.5, 0.55, desc, size=15, color=INK, anchor=MSO_ANCHOR.MIDDLE)
button(s, 0.7, 6.35, 11.9, 'Срок не опубликован? Не угадываем — пользователь указывает дату сам, когда её объявят', h=0.55)
notes(s, 'Это прямой ответ на принцип кейса «Evidence over assumptions»: мы показываем пользователю, насколько можно доверять каждому факту.')

# ===================================================================== 9. Конкуренты
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
brand(s)
heading(s, 'КОНКУРЕНТЫ', 'ЧЕМ МЫ ОТЛИЧАЕМСЯ')
rows = [('Telegram-каналы с объявлениями', 'Лента постов: быстро, но без проверки условий и без плана; старое уходит вниз.'),
        ('Сайты-агрегаторы возможностей', 'Большие каталоги на английском; фильтры по стране и типу, но не по твоему профилю.'),
        ('Официальные сайты программ', 'Точные условия, но у каждой программы свой сайт, свой язык и свой календарь.'),
        ('Школьный учитель / профориентатор', 'Знает учеников, но не успевает следить за десятками программ и сроков.')]
text(s, 5.35, 2.4, 7.2, 0.3, 'ЧТО ЕСТЬ СЕЙЧАС', size=11, color=MUTED)
for i, (n, what) in enumerate(rows):
    y = 2.75 + i * 0.78
    shape(s, MSO_SHAPE.RECTANGLE, 0.7, y, 11.9, 0.68, BG, LINE, 1)
    text(s, 0.95, y + 0.08, 4.2, 0.55, n, size=14, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 5.35, y + 0.06, 7.1, 0.58, what, size=12.5, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
button(s, 0.7, 6.0, 11.9, 'Maqsat: проверка условий под твой профиль + план назад от дедлайна + календарь + честные данные', h=0.6)
notes(s, 'Не говорите, что аналогов нет: жюри знает Telegram-каналы и агрегаторы. Наше отличие — персональная проверка условий и план.')

# ===================================================================== 10. Техника и UX
s = prs.slides.add_slide(BLANK)
bg(s, BG)
brand(s)
heading(s, 'ТЕХНИКА', 'ПРОСТО, НАДЁЖНО, ДОСТУПНО', w=8.0)
picture(s, IMG / 'mobile.png', 9.95, 1.2, 2.65, crop_bottom=0.1)
cols = [('Архитектура', ['HTML/CSS/JS без зависимостей', 'Данные отделены от движка: новая возможность — одна запись', 'Календарь по стандарту iCalendar (RFC 5545)', '42 автотеста — запускаются на GitHub при каждом изменении', 'Раз в неделю робот проверяет все ссылки-источники']),
        ('Доступность', ['Без регистрации; данные — только на устройстве', 'Работает офлайн (PWA), ставится на телефон', 'Русский и қазақша', 'Клавиатура, экранные дикторы, крупные кнопки', 'Тёмная тема, печать для родителей'])]
for i, (h_, items) in enumerate(cols):
    x = 0.7 + i * 4.65
    shape(s, MSO_SHAPE.RECTANGLE, x, 2.55, 4.4, 3.95, WHITE, LINE, 1)
    text(s, x + 0.3, 2.8, 3.8, 0.45, h_, size=19, color=BLUE)
    tb = text(s, x + 0.3, 3.4, 3.9, 3.0, items, size=14, color=GRAY, spacing=8)
    bullets(tb)
footer(s)

# ===================================================================== 11. Эффект и дальше
s = prs.slides.add_slide(BLANK)
s.shapes.add_picture(str(TUNNEL), 0, 0, width=prs.slide_width, height=prs.slide_height)
brand(s)
dots(s, [(0.55, 0.45, 0.2), (1.35, 0.9, 0.16), (2.3, 0.45, 0.2)])
heading(s, 'ДАЛЬШЕ', 'КАК ИЗМЕРИМ ЭФФЕКТ', y=1.2)
metrics = [('Заявки до дедлайна', 'главная метрика'), ('Шаги в срок', 'работает ли план'), ('< 3 минут', 'до готового плана'), ('«Почти» → «подходишь»', 'помогает ли разбор')]
for i, (h_, b) in enumerate(metrics):
    x = 0.8 + i * 3.1
    target(s, x + 0.3, 3.55, 0.55)
    text(s, x + 0.75, 3.3, 2.3, 0.5, h_, size=15, color=BLUE)
    text(s, x + 0.75, 3.75, 2.3, 0.4, b, size=12, color=GRAY)
road = 'Вычитка казахского и перевод описаний · база 100+ возможностей с проверкой модератором · сравнение содержимого страниц-источников · Telegram-напоминания · режим для профориентатора'
text(s, 0.8, 4.6, 11.8, 0.8, [[('Дорожная карта: ', {'color': BLUE}), (road, {'color': INK})]], size=14)
text(s, 0.8, 5.85, 8.6, 0.4, 'Команда: Bolatbay Yersultan · Zharkynuly Eren · Kydyrkhan Olzhas', size=16, color=INK)
text(s, 0.8, 6.3, 8.6, 0.35, 'GitHub: github.com/sas1k00/hackatonishe   ·   Демо: sas1k00.github.io/hackatonishe', size=13, color=GRAY)
button(s, 10.15, 5.95, 2.5, 'Спасибо!')

prs.save(OUT)
print('saved', OUT)
